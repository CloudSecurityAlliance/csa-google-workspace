"""What the register does when part of the document will not answer.

`export_comments` assembles a table out of several independent reads: the comments, the cell
grids behind them, the headers around those cells, the notes count. Any of those can fail on
its own - a tab this account cannot read, a Sheets API that is off, a file type with no notes
at all - and none of them is a reason to lose the register.

So each degrades to a blank rather than an exception, and the caveats say which blanks are
"nothing there" and which are "we did not look". That distinction is the whole value of the
artifact: a reviewer reading a blank `row_header` must not conclude the cell has no header.
"""
from __future__ import annotations

import pytest

from csa_google_workspace import _export


class Document:
    """A spreadsheet-shaped document whose parts can be made to fail individually."""

    def __init__(self, *, tabs=("Sheet1",), grid=None, values_raises=False,
                 notes=None, notes_raises=False):
        # `tabs` is a property of plain title STRINGS on a Sheet, and `notes` an iterable
        # attribute rather than a method - both are read as attributes by `_export`, so the
        # double has to present them that way or it tests a shape nothing produces.
        self.tabs = list(tabs)
        self._grid = grid if grid is not None else [["Header", "Other"], ["A2", "B2"]]
        self._values_raises = values_raises
        if notes_raises:
            self.notes = _RaisingIterable()
        elif notes is not None:
            self.notes = list(notes)

    def values(self, tab):
        if self._values_raises:
            raise RuntimeError("that tab is not readable by this account")
        return self._grid


class _RaisingIterable:
    """Iterating it fails - the shape a notes lookup takes when Sheets is unavailable."""

    def __iter__(self):
        raise RuntimeError("the Sheets API is not enabled on this project")


class TestATabThatWillNotBeRead:
    def test_its_grid_is_empty_and_the_others_are_not_lost(self):
        """A file shared with this account can still contain a tab it cannot read - protected
        ranges, a partial share. Failing the export would lose the comments on every OTHER
        tab, which are the ones somebody asked for."""
        grids, tabs = _export._cell_lookup(Document(values_raises=True))

        assert tabs == ["Sheet1"], "the tab is still listed"
        assert grids == {"Sheet1": []}, "and its grid is empty rather than absent"

    def test_a_readable_tab_comes_back_whole(self):
        grids, _ = _export._cell_lookup(Document())
        assert grids["Sheet1"] == [["Header", "Other"], ["A2", "B2"]]

    def test_a_document_with_no_tabs_is_not_an_error(self):
        """A Doc or a deck. `_cell_lookup` is asked unconditionally and has to answer for a
        file type that has no grid at all."""
        assert _export._cell_lookup(Document(tabs=())) == ({}, [])


class TestCountingNotes:
    def test_a_file_type_with_no_notes_counts_zero(self):
        """`notes` is absent on a Doc. Zero, not an exception - and the caveat that mentions
        notes then correctly does not fire."""

        class NoNotes:
            tabs = []

        assert _export._note_count(NoNotes()) == 0

    def test_notes_that_cannot_be_read_count_zero_rather_than_failing(self):
        """The Sheets API being off is the ordinary cause. A count that is only a caveat must
        not be able to take down the register it is a caveat on."""
        assert _export._note_count(Document(notes_raises=True)) == 0

    def test_notes_that_can_be_read_are_counted(self):
        assert _export._note_count(Document(notes=[1, 2, 3])) == 3


class TestFlatteningAValueForACell:
    """Everything in the register is text by the time it reaches CSV or a sheet, and each
    conversion is a decision about what somebody will read."""

    def test_none_becomes_empty_rather_than_the_word_none(self):
        """A literal "None" in a cell reads as data. Half a review register would say it."""
        assert _export.flatten(None) == ""

    @pytest.mark.parametrize("value, expected", [(True, "yes"), (False, "no")])
    def test_booleans_become_words_not_TRUE_and_FALSE(self, value, expected):
        """The register's own action columns use TRUE/FALSE, and a DATA column rendering the
        same two words would look like something to fill in."""
        assert _export.flatten(value) == expected

    def test_a_dict_is_flattened_to_readable_pairs(self):
        assert _export.flatten({"a": 1, "b": 2}) == "a=1 | b=2"

    def test_anything_else_is_its_string(self):
        assert _export.flatten(42) == "42"


class TestDroppingColumnsNothingFilled:
    def test_an_empty_register_keeps_every_column(self):
        """Nothing to measure emptiness against. A header of nothing would make the file
        unopenable as a register, and an export with no comments is an ordinary result."""
        columns = ["thread_id", "cell", "row_header"]
        assert _export.used_columns(columns, []) == columns

    def test_a_column_nothing_filled_is_dropped(self):
        assert "row_header" not in _export.used_columns(
            ["thread_id", "row_header"], [{"thread_id": "t1", "row_header": ""}])

    def test_a_column_something_filled_is_kept(self):
        assert "row_header" in _export.used_columns(
            ["thread_id", "row_header"], [{"thread_id": "t1", "row_header": "Revenue"}])


class TestWhereAGivenPathLands:
    def test_a_path_with_a_separator_is_taken_as_given(self, tmp_path):
        """Somebody who typed a path meant that path. Joining it onto the export directory
        would put the file somewhere neither of them named."""
        target, note = _export.resolve_export_path(
            str(tmp_path / "out.csv"), default_dir=str(tmp_path / "elsewhere"),
            doc_name="d", stamp="20260927", suffix=".csv")

        assert target == tmp_path / "out.csv"
        assert not note, "nothing to explain: it went where they said"

    def test_a_tilde_path_is_expanded_and_not_joined(self, tmp_path):
        home = tmp_path / "home"
        home.mkdir()
        target, _ = _export.resolve_export_path(
            "~/out.csv", default_dir=str(tmp_path), doc_name="d", stamp="s", suffix=".csv")
        assert target.is_absolute() and target.name == "out.csv"

    def test_a_bare_name_lands_in_the_export_directory_and_says_so(self, tmp_path):
        """A name with no separator is a name, not a path. The note exists because the file
        then appears somewhere the caller did not type."""
        target, note = _export.resolve_export_path(
            "out.csv", default_dir=str(tmp_path), doc_name="d", stamp="s", suffix=".csv")

        assert target == tmp_path / "out.csv"
        assert note, "a file written somewhere the caller did not name must say where"

    def test_a_bare_name_with_no_export_directory_is_refused(self, tmp_path):
        """Rather than guessing at the working directory, which for a stdio server is wherever
        the MCP client happened to be launched from."""
        with pytest.raises(ValueError, match="no default export directory"):
            _export.resolve_export_path("out.csv", default_dir="", doc_name="d",
                                        stamp="s", suffix=".csv")


class TestTheXlsxExtra:
    def test_it_reports_what_is_installed(self):
        """True here, because `openpyxl` ships with the `mcp` extra and the test environment
        has it. The value is not the point - reaching the check is."""
        assert _export.xlsx_supported() in (True, False)

    def test_a_missing_openpyxl_is_false_rather_than_an_import_error(self, monkeypatch):
        """A library-only embedder may not have it, and a tool that used to work should
        degrade to the unformatted path rather than fail. Asked BEFORE anything is built, so
        the answer arrives before the work."""
        import builtins

        real_import = builtins.__import__

        def refuse(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("No module named 'openpyxl'")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", refuse)
        assert _export.xlsx_supported() is False


class Location:
    def __init__(self, cell, tab, row, col):
        self.cell, self.tab, self.row, self.col = cell, tab, row, col


class Comment:
    def __init__(self, comment_id, *, location=None, content="A question"):
        self.id, self.location, self.content = comment_id, location, content
        self.author = None
        self.resolved = False
        self.replies = []
        self.quoted_text = None


class TestTheHeaderGuessIsDeclaredAsAGuess:
    """`row_header` and `column_header` are inferred from column A and row 1 - the usual
    layout of a spreadsheet, and not a guaranteed one.

    The caveat is the whole mechanism. A register naming "Revenue / Q3" as the meaning of a
    cell is what somebody quotes in a summary, and on a sheet with a title block above the
    table, a transposed layout, or merged headers, that label is wrong while looking exactly
    as confident as a right one. So the columns are offered AND the guess is declared, in the
    same artifact.
    """

    GRID = [["", "Q1", "Q2", "Q3"],
            ["Revenue", "10", "20", "30"],
            ["Costs", "1", "2", "3"]]

    def document(self):
        return Document(tabs=("Sheet1",), grid=self.GRID)

    def test_a_cell_anchored_comment_gets_both_headers(self):
        """D2, as 1-BASED row 2 and column 4 - the same convention as `_at` and
        `Location.row`/`col`. `_headers` was written 0-based first and silently returned no
        headers at all rather than failing, and this test's first draft made the same mistake
        from the other side: off-by-one here looks like "the feature does not work"."""
        comment = Comment("t1", location=Location("D2", "Sheet1", 2, 4))
        _, rows, _ = _export.comment_rows(self.document(), [comment])

        assert rows[0]["row_header"] == "Revenue"
        assert rows[0]["column_header"] == "Q3"
        assert rows[0]["cell_text"] == "30"

    def test_the_caveat_fires_when_a_header_was_guessed(self):
        comment = Comment("t1", location=Location("D2", "Sheet1", 2, 4))
        _, _, caveats = _export.comment_rows(self.document(), [comment])

        guess = [c for c in caveats if "GUESS" in c]
        assert guess, "a header was inferred and the register did not say so"
        assert "cell_text" in guess[0], "the caveat must name what to check it against"

    def test_it_does_not_fire_when_nothing_was_guessed(self):
        """A file-level comment on the same sheet. A caveat that fired regardless would be one
        more sentence to skip on every register, which is how the real ones get skipped too."""
        _, _, caveats = _export.comment_rows(self.document(), [Comment("t1")])
        assert not [c for c in caveats if "GUESS" in c]


class TestTheWorkbookWriterAndEmptyCells:
    def test_an_empty_value_does_not_break_the_text_forcing_pass(self):
        """Every data cell is forced to text, because openpyxl INFERS type from value and a
        comment body beginning `=` was being written as a live formula element (#182). An
        empty cell has `None` rather than `""`, so it is not a string and is skipped - the
        loop has to tolerate that rather than assume every cell it walks is one."""
        pytest.importorskip("openpyxl")
        import io

        import openpyxl

        data = _export.to_xlsx_bytes(
            ["thread_id", "text"],
            [{"thread_id": "t1", "text": ""}, {"thread_id": "t2", "text": "=IMPORTXML(1,2)"}],
            title="Comments")
        ws = openpyxl.load_workbook(io.BytesIO(data)).active

        assert ws.cell(row=2, column=2).value in (None, "")
        assert ws.cell(row=3, column=2).value == "=IMPORTXML(1,2)"
        assert ws.cell(row=3, column=2).data_type == "s", "a formula was written as one"
